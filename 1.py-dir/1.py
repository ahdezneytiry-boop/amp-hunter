from Bio import Entrez
from docx import Document

Entrez.email = "ahdezneytiry@gmail.com"

def pubmed_abstracts():
    res = Entrez.esearch(db="pubmed", term="antimicrobial peptide[Title]", retmax=5)
    ids = Entrez.read(res)["IdList"]
    res.close()

    rec = Entrez.efetch(db="pubmed", id=",".join(ids), retmode="xml")
    articles = Entrez.read(rec)["PubmedArticle"]
    rec.close()

    doc = Document()
    doc.add_heading("PubMed Abstracts", 0)
    for p in articles:
        art = p["MedlineCitation"]["Article"]
        title = art.get("ArticleTitle", "无标题")
        text = "无摘要"
        if art.get("Abstract") and art["Abstract"].get("AbstractText"):
            absdata = art["Abstract"]["AbstractText"]
            if isinstance(absdata, list):
                text = " ".join([str(x) for x in absdata])
            else:
                text = str(absdata)
        doc.add_heading(title, level=1)
        doc.add_paragraph(text)
    doc.save("PubMed_Abstracts1.docx")

if __name__ == "__main__":
    pubmed_abstracts()