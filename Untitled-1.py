
import time
from Bio import Entrez
from docx import Document

def pubmed_fetch_abstracts(query, retmax=200):
    Entrez.email = "ahdezneytiry@gmail.com"
    # 1. ESearch: 找到相关 PMID
    handle = Entrez.esearch(db="pubmed",
                            term=query + "[Title]",
                            retmax=retmax,
                            sort="relevance")
    rec = Entrez.read(handle)
    handle.close()
    pmids = rec["IdList"]
    if not pmids:
        return []

    # 2. EFetch: 拉摘要
    handle = Entrez.efetch(db="pubmed",
                           id=",".join(pmids),
                           rettype="abstract",
                           retmode="text")
    data = handle.read()
    handle.close()

    # 如果要结构化解析每篇，可以用 EFetch XML，但这里按文本方式简单截取
    # 下面用 XML 方式更好
    handle = Entrez.efetch(db="pubmed",
                           id=",".join(pmids),
                           rettype="xml",
                           retmode="xml")
    records = Entrez.read(handle)
    handle.close()

    items = []
    for rec in records:
        title = ""
        abstract = ""
        if "ArticleTitle" in rec["MedlineCitation"]["Article"]:
            title = rec["MedlineCitation"]["Article"]["ArticleTitle"]
        ab = rec["MedlineCitation"]["Article"].get("Abstract")
        if ab and "AbstractText" in ab:
            if isinstance(ab["AbstractText"], list):
                abstract = "\n".join(str(x) for x in ab["AbstractText"])
            else:
                abstract = str(ab["AbstractText"])
        items.append((title, abstract))
    return items


def save_to_word(items, filename):
    doc = Document()
    doc.add_heading("PubMed 抗菌肽（antimicrobial peptide）标题文章摘要", level=1)
    for idx, (title, abstract) in enumerate(items, 1):
        doc.add_heading(f"{idx}. {title}", level=2)
        if abstract:
            doc.add_paragraph(abstract)
        else:
            doc.add_paragraph("（无摘要）")
        doc.add_page_break()
    doc.save(filename)

if __name__ == "__main__":
    query_phrase = "antimicrobial peptide"
    print("开始检索 PubMed ...")
    articles = pubmed_fetch_abstracts(query_phrase, retmax=100)  # 100 可改
    print(f"查到 {len(articles)} 条（含标题关键词）的记录。")
    out = "pubmed_antimicrobial_peptide_abstracts.docx"
    save_to_word(articles, out)
    print("保存完成：", out)